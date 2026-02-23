# Monitoring/reporting/report_exporter.py
# ============================================================
# A4 PDF Export (Full-page layout) + auto-fit + auto-split tall images
# - Fix LayoutError (image too tall)
# - Make images fill A4 page as much as possible
# - Optional: split very tall images into multiple pages (recommended)
#
# Requirements:
#   py -3.13 -m pip install reportlab pillow python-docx
# ============================================================

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from typing import List, Optional, Tuple

# ===== ReportLab =====
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image as RLImage,
    PageBreak,
    Table,
    TableStyle,
    KeepInFrame,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

# ===== Pillow (for splitting tall images) =====
try:
    from PIL import Image as PILImage
except Exception:
    PILImage = None


# ============================================================
# Config
# ============================================================

@dataclass
class PdfLayoutConfig:
    pagesize: Tuple[float, float] = A4

    # Margins: reduce to make content fill the page
    margin_left: float = 10 * mm
    margin_right: float = 10 * mm
    margin_top: float = 10 * mm
    margin_bottom: float = 10 * mm

    # Header/Footer
    show_header: bool = True
    show_footer: bool = True
    header_height: float = 10 * mm
    footer_height: float = 10 * mm

    # Title section sizes
    title_on_first_page_only: bool = True
    title_font_size: int = 20
    section_font_size: int = 13

    # Image behavior
    split_tall_images: bool = True     # best readability
    always_full_width: bool = True     # scale by width first
    max_split_part_px: int = 0         # 0 = auto compute from A4 frame
    image_padding_top: float = 2
    image_padding_bottom: float = 2

    # Optional light border around images
    image_border: bool = False


# ============================================================
# Utilities
# ============================================================

def _ensure_dir(path: str) -> None:
    os.makedirs(path, exist_ok=True)


def _temp_dir(subfolder: str = "dashboard_reports") -> str:
    base = os.path.join(tempfile.gettempdir(), subfolder)
    _ensure_dir(base)
    return base


def _get_usable_area(cfg: PdfLayoutConfig) -> Tuple[float, float]:
    page_w, page_h = cfg.pagesize
    usable_w = page_w - cfg.margin_left - cfg.margin_right
    usable_h = page_h - cfg.margin_top - cfg.margin_bottom

    # Reserve header/footer space inside the frame (so images don't collide)
    if cfg.show_header:
        usable_h -= cfg.header_height
    if cfg.show_footer:
        usable_h -= cfg.footer_height

    return usable_w, usable_h


def _draw_header_footer(canvas, doc, title: str, cfg: PdfLayoutConfig):
    page_w, page_h = cfg.pagesize
    canvas.saveState()

    if cfg.show_header:
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(cfg.margin_left, page_h - cfg.margin_top + 2, title)

    if cfg.show_footer:
        canvas.setFont("Helvetica", 9)
        canvas.drawRightString(
            page_w - cfg.margin_right,
            cfg.margin_bottom - 8,
            f"Page {doc.page}"
        )

    canvas.restoreState()


def _fit_image(img_path: str, max_w: float, max_h: float, prefer_full_width: bool = True) -> RLImage:
    """
    Create RLImage and set drawWidth/drawHeight so it fits within max_w/max_h.
    prefer_full_width=True: try to fill width as much as possible.
    """
    img = RLImage(img_path)
    iw, ih = float(img.imageWidth), float(img.imageHeight)
    if iw <= 0 or ih <= 0:
        return img

    if prefer_full_width:
        # Scale to full width first, then clamp by height if needed
        scale = max_w / iw
        new_h = ih * scale
        if new_h > max_h:
            # too tall -> scale by height instead
            scale = max_h / ih
    else:
        scale = min(max_w / iw, max_h / ih)

    scale = min(scale, 1.0e9)  # just in case
    img.drawWidth = iw * scale
    img.drawHeight = ih * scale
    return img


def _keep_in_frame(flowable, max_w: float, max_h: float):
    """
    Extra safety: shrink if slightly overflowing.
    """
    return KeepInFrame(max_w, max_h, [flowable], mode="shrink")


def _estimate_split_height_px(img_path: str, max_w_pt: float, max_h_pt: float) -> int:
    """
    Compute a safe per-page slice height (pixels) so each slice fits when scaled.
    """
    if PILImage is None:
        return 1200

    im = PILImage.open(img_path)
    w_px, h_px = im.size
    if w_px <= 0:
        return 1200

    # If we scale by width to max_w_pt:
    # scale = max_w_pt / w_px
    # slice_h_pt = slice_h_px * scale <= max_h_pt
    # => slice_h_px <= max_h_pt * w_px / max_w_pt
    slice_h_px = int(max_h_pt * w_px / max_w_pt)

    # Safety margin for any paddings
    slice_h_px = int(slice_h_px * 0.95)

    # clamp
    return max(400, min(slice_h_px, 5000))


def _split_image(img_path: str, out_dir: str, part_height_px: int, prefix: str) -> List[str]:
    """
    Split tall image into multiple PNG parts.
    """
    if PILImage is None:
        return [img_path]

    im = PILImage.open(img_path)
    w, h = im.size
    if h <= part_height_px:
        return [img_path]

    parts: List[str] = []
    y = 0
    idx = 1
    while y < h:
        box = (0, y, w, min(y + part_height_px, h))
        part = im.crop(box)
        out = os.path.join(out_dir, f"{prefix}_{idx}.png")
        part.save(out)
        parts.append(out)
        y += part_height_px
        idx += 1

    return parts


def _make_image_block(
    img_path: str,
    usable_w: float,
    usable_h: float,
    cfg: PdfLayoutConfig,
) -> Table:
    """
    Wrap image in a 1x1 table for consistent centering/padding (no random shrinking).
    """
    img = _fit_image(
        img_path,
        max_w=usable_w,
        max_h=usable_h,
        prefer_full_width=cfg.always_full_width,
    )

    # final safety net
    kif = _keep_in_frame(img, usable_w, usable_h)

    tbl = Table([[kif]], colWidths=[usable_w])
    style_cmds = [
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), cfg.image_padding_top),
        ("BOTTOMPADDING", (0, 0), (-1, -1), cfg.image_padding_bottom),
    ]
    if cfg.image_border:
        style_cmds.append(("BOX", (0, 0), (-1, -1), 0.35, colors.lightgrey))

    tbl.setStyle(TableStyle(style_cmds))
    return tbl


def _collect_cell_images_from_temp(temp_folder: str) -> List[str]:
    """
    Collect images named like cell_r1_c1.png in temp folder.
    """
    if not os.path.isdir(temp_folder):
        return []
    out = []
    for fn in sorted(os.listdir(temp_folder)):
        low = fn.lower()
        if low.endswith(".png") and low.startswith("cell_"):
            out.append(os.path.join(temp_folder, fn))
    return out


# ============================================================
# Public API
# ============================================================

def export_report_pdf(
    save_path: str,
    title: str,
    parent_widget=None,
    *,
    temp_image_dir: Optional[str] = None,
    image_paths: Optional[List[str]] = None,
    cfg: Optional[PdfLayoutConfig] = None,
) -> None:
    """
    Export report PDF in A4, full-page layout.

    How to use (recommended):
    - If your app already exports images (cell_rX_cY.png) to temp folder,
      call export_report_pdf(save_path, title, self, temp_image_dir=that_dir)

    - Or pass explicit list of image_paths.

    Fixes:
    - Auto-fit full width to A4
    - Auto-split tall images into multiple pages (prevents LayoutError)
    """
    cfg = cfg or PdfLayoutConfig()
    temp_image_dir = temp_image_dir or _temp_dir("dashboard_reports")

    # If no image_paths passed, auto collect from temp folder
    if image_paths is None:
        image_paths = _collect_cell_images_from_temp(temp_image_dir)

    doc = SimpleDocTemplate(
        save_path,
        pagesize=cfg.pagesize,
        leftMargin=cfg.margin_left,
        rightMargin=cfg.margin_right,
        topMargin=cfg.margin_top,
        bottomMargin=cfg.margin_bottom,
        title=title,
        author="Dashboard-VP",
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="BigTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=cfg.title_font_size,
        spaceAfter=10,
    ))
    styles.add(ParagraphStyle(
        name="Section",
        parent=styles["Heading2"],
        fontSize=cfg.section_font_size,
        spaceBefore=6,
        spaceAfter=6,
    ))

    usable_w, usable_h = _get_usable_area(cfg)

    story: List = []

    # ===== Page 1: Title + optional summary =====
    story.append(Paragraph(title, styles["BigTitle"]))
    story.append(Paragraph("Monitoring Layout", styles["Section"]))
    story.append(Spacer(1, 4))

    # If you want the FIRST image to fill the first page too:
    # set cfg.title_on_first_page_only=False and reduce title sizes
    # But the best look is: title page then full-page images.
    if cfg.title_on_first_page_only and image_paths:
        story.append(PageBreak())

    # ===== Pages: Images =====
    for img_idx, img_path in enumerate(image_paths, start=1):
        if not os.path.exists(img_path):
            continue

        # For a single image: optionally split to multiple pages
        if cfg.split_tall_images:
            prefix = os.path.splitext(os.path.basename(img_path))[0] + "_part"

            if cfg.max_split_part_px > 0:
                part_h_px = cfg.max_split_part_px
            else:
                part_h_px = _estimate_split_height_px(img_path, usable_w, usable_h)

            parts = _split_image(img_path, temp_image_dir, part_h_px, prefix)
        else:
            parts = [img_path]

        # Optional section label for each block (small)
        # story.append(Paragraph(os.path.basename(img_path), styles["Section"]))
        # story.append(Spacer(1, 2))

        for p_i, part in enumerate(parts, start=1):
            # On the very first image page (if not using title-only page),
            # usable_h is slightly reduced by title text, but we keep it simple:
            # If you want perfect: compute reserved space and reduce max_h.
            block = _make_image_block(part, usable_w, usable_h, cfg)
            story.append(block)

            # If more parts exist, force new page
            if p_i < len(parts):
                story.append(PageBreak())

        # After each original image, start new page (recommended)
        # to get true "full page" look per chart/layout
        if img_idx < len(image_paths):
            story.append(PageBreak())

    # If no images, still build PDF
    if not image_paths:
        story.append(Paragraph("No exported images found to include in this report.", styles["BodyText"]))

    doc.build(
        story,
        onFirstPage=lambda c, d: _draw_header_footer(c, d, title, cfg),
        onLaterPages=lambda c, d: _draw_header_footer(c, d, title, cfg),
    )


def export_report_docx(save_path: str, title: str, parent_widget=None) -> None:
    """
    Basic DOCX exporter (safe fallback).
    If your project already has a richer docx exporter, keep yours.
    """
    try:
        from docx import Document
    except Exception as e:
        raise RuntimeError("python-docx is required. Install: pip install python-docx") from e

    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph("Auto-generated report.")
    doc.save(save_path)


# ============================================================
# Optional: quick CLI test (run this file directly)
# ============================================================
if __name__ == "__main__":
    # Put your temp folder here (where cell_r1_c1.png exists)
    tmp = _temp_dir("dashboard_reports")
    out_pdf = os.path.join(tmp, "test_monitoring_report_A4.pdf")

    export_report_pdf(
        out_pdf,
        "Monitoring Report",
        temp_image_dir=tmp,
        cfg=PdfLayoutConfig(
            split_tall_images=True,
            title_on_first_page_only=True,
            always_full_width=True,
            margin_left=8 * mm,
            margin_right=8 * mm,
            margin_top=8 * mm,
            margin_bottom=8 * mm,
            image_border=False,
        ),
    )
    print("Exported:", out_pdf)